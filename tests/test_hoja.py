"""
Formato editable de la hoja de revisión: la plantilla de Word con marcadores y el
archivo de frases. El revisor los cambia sin tocar el programa.
"""
import copy

import pytest
from docx import Document

from reportes.reporte_resumen import (FRASES, cargar_frases, crear_plantilla_hoja, escribir_resumen,
                                      lista_de_lineas, redactar)

DATOS = dict(fecha_revision="27/09/2026", n_revision="1", proyecto="Arquitectura de datos",
             archivo="tesis.docx", asesor="Pari C.", revisor="kevin", tesista="", codigo="", similitud="")


def _texto(ruta):
    return [p.text for p in Document(ruta).paragraphs if p.text.strip()]


@pytest.mark.parametrize("numeros, esperado", [
    ([], ""),
    ([33], "Línea 33"),
    ([126, 11], "Líneas 11 y 126"),
    ([11, 126, 633, 11], "Líneas 11, 126 y 633"),
])
def test_lista_de_lineas(numeros, esperado):
    assert lista_de_lineas(numeros) == esperado


def test_con_muchas_lineas_se_muestran_las_primeras_y_otros():
    frases = copy.deepcopy(FRASES)
    frases["lineas"]["maximo"] = 3
    assert lista_de_lineas([1, 2, 3, 4, 5], frases) == "Líneas 1, 2, 3 y otros"


def test_la_plantilla_del_programa_da_la_hoja_de_siempre(tmp_path):
    ruta = tmp_path / "hoja.docx"
    items = [dict(categoria="Ortografía", severidad="Advertencia", observacion="Doble espacio",
                  ubicacion="x", nlinea=40, nlineas=[40, 90])]
    escribir_resumen(str(ruta), DATOS, redactar(items, []))
    texto = _texto(ruta)
    assert texto[0] == "REVISIÓN PGI 27/09/2026 – 1"
    assert "Asesor: Pari C." in texto
    assert not any("Tesista" in t for t in texto), "{{TESISTA?}} vacío borra su línea"
    assert "Líneas 40 y 90: corregir ortografía y gramática" in texto
    assert not any("{{" in t for t in texto)
    obs = next(p for p in Document(ruta).paragraphs if p.text.startswith("Líneas 40"))
    assert obs.runs[0].bold and obs.runs[0].text == "Líneas 40 y 90: ", "el prefijo de líneas va en negrita"


def test_sin_numero_de_revision_no_queda_el_guion_colgando(tmp_path):
    ruta = tmp_path / "hoja.docx"
    escribir_resumen(str(ruta), dict(DATOS, n_revision=""), [])
    assert _texto(ruta)[0] == "REVISIÓN PGI 27/09/2026"


def test_una_plantilla_editada_en_word_se_llena(tmp_path):
    """Word parte los marcadores en varios fragmentos al editar: igual se reemplazan"""
    plantilla = tmp_path / "propia.docx"
    doc = Document()
    p = doc.add_paragraph()
    for trozo in ("HOJA DE REVISIÓN · Revisor: {{RE", "VIS", "OR}} · {{FECHA_", "REVISION}}"):
        p.add_run(trozo)
    doc.add_paragraph("Código: {{CODIGO?}}")
    doc.add_paragraph("{{OBSERVACIONES}}")
    doc.add_paragraph("Firma")
    doc.save(str(plantilla))
    salida = tmp_path / "hoja.docx"
    escribir_resumen(str(salida), DATOS, redactar([], [dict(nlinea=11, nlineas=[11])]), str(plantilla))
    assert _texto(salida) == ["HOJA DE REVISIÓN · Revisor: kevin · 27/09/2026",
                              "Revisión gramatical y ortográfica de todo el proyecto",
                              "Línea 11: corregir ortografía y gramática", "Firma"]


def test_si_la_plantilla_perdio_el_marcador_las_observaciones_van_al_final(tmp_path):
    plantilla = tmp_path / "sin_marcador.docx"
    doc = Document()
    doc.add_paragraph("Encabezado")
    doc.save(str(plantilla))
    salida = tmp_path / "hoja.docx"
    escribir_resumen(str(salida), DATOS, redactar([], [dict(nlinea=5, nlineas=[5])]), str(plantilla))
    assert _texto(salida)[-1] == "Línea 5: corregir ortografía y gramática"


def test_las_frases_del_revisor_reemplazan_las_del_programa(tmp_path):
    ruta = tmp_path / "frases.yaml"
    ruta.write_text('lineas:\n  varias: "Ver líneas {lista}"\nobservaciones:\n'
                    '  ortografia: "revisar la ortografía"\n', encoding="utf-8")
    frases = cargar_frases(str(ruta))
    assert frases["observaciones"]["referencias"] == FRASES["observaciones"]["referencias"], \
        "lo que el revisor no escribió queda como viene"
    hoja = redactar([], [dict(nlinea=1, nlineas=[1]), dict(nlinea=9, nlineas=[9])], frases=frases)
    assert "Ver líneas 1 y 9: revisar la ortografía" in hoja


def test_un_archivo_de_frases_con_error_avisa_y_usa_las_originales(tmp_path, monkeypatch):
    from core import orquestador
    carpeta = tmp_path / "reporte"
    carpeta.mkdir()
    (carpeta / "frases_reporte.yaml").write_text("lineas: [sin cerrar\n", encoding="utf-8")
    monkeypatch.setattr(orquestador, "DATOS", str(tmp_path))
    avisos = []
    _, frases = orquestador.formato_de_hoja(avisos.append)
    assert frases == FRASES
    assert any("error de formato" in a for a in avisos)


def test_el_yaml_que_trae_el_programa_coincide_con_los_valores_por_defecto():
    """si alguien cambia uno sin el otro, 'Restablecer' dejaría un formato distinto"""
    from core.config import ruta_recurso
    assert cargar_frases(ruta_recurso("reporte/frases_reporte.yaml")) == FRASES


def test_la_plantilla_nueva_tiene_todos_los_marcadores(tmp_path):
    texto = " ".join(p.text for p in crear_plantilla_hoja().paragraphs)
    for marcador in ("{{FECHA_REVISION}}", "{{TITULO}}", "{{ASESOR}}", "{{OBSERVACIONES}}", "{{REVISOR}}"):
        assert marcador in texto
