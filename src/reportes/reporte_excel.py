"""
Reporte en Excel: una hoja con todas las observaciones, otra con el resumen por
categoría y otra con el detalle de ortografía. Es la salida opcional (--excel),
pensada para filtrar y ordenar; el entregable formal es el Word.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


COLORES = {"Error": "F8CBAD", "Advertencia": "FFE699", "Revisar": "BDD7EE"}


def escribir_reporte(ruta_salida, nombre_doc, obs, filas_orto, motor):
    wb = Workbook()
    ws = wb.active
    ws.title = "Observaciones"
    ws.append([f"Revisión automática: {nombre_doc}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Borrador generado por el revisor. Validar cada observación antes de enviarla al tesista."])
    ws.append([])
    enc = ["N°", "Categoría", "Severidad", "Página y línea", "Ubicación", "Observación", "Detalle / corrección"]
    ws.append(enc)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    orden_sev = {"Error": 0, "Advertencia": 1, "Revisar": 2}
    orden_cat = {"Estructura": 0, "Formato": 1, "Extensión": 2, "Citas": 3, "Ortografía": 4}
    items = sorted(obs.items, key=lambda x: (orden_cat.get(x["categoria"], 9), orden_sev.get(x["severidad"], 9)))
    for k, it in enumerate(items, 1):
        ws.append([k, it["categoria"], it["severidad"], it.get("linea", ""), it["ubicacion"], it["observacion"], it["detalle"]])
        ws.cell(row=ws.max_row, column=3).fill = PatternFill("solid", fgColor=COLORES.get(it["severidad"], "FFFFFF"))
    for col, ancho in zip("ABCDEFG", [5, 13, 12, 18, 40, 55, 60]):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=5):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"

    r = wb.create_sheet("Resumen")
    r.append(["Categoría", "Errores", "Advertencias", "Revisar"])
    for c in r[1]:
        c.font = Font(bold=True)
    for cat in orden_cat:
        sub = [x for x in obs.items if x["categoria"] == cat]
        r.append([cat] + [sum(1 for x in sub if x["severidad"] == s) for s in ("Error", "Advertencia", "Revisar")])
    r.append([])
    r.append(["Motor de ortografía", motor])
    r.column_dimensions["A"].width = 22
    r.column_dimensions["B"].width = 14

    o = wb.create_sheet("Ortografía")
    o.append(["Página y línea", "Ubicación", "Palabra", "Problema", "Sugerencias", "Contexto"])
    for c in o[1]:
        c.font = Font(bold=True)
    for f in filas_orto:
        o.append([f.get("linea", ""), f["parrafo"], f["palabra"], f["problema"], f["sugerencias"], f["contexto"]])
    for col, ancho in zip("ABCDEF", [18, 40, 20, 45, 30, 70]):
        o.column_dimensions[col].width = ancho
    wb.save(ruta_salida)
