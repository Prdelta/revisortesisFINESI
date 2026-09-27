"""
Documentos con citas de Mendeley o Zotero.

Estos gestores guardan cada cita, y la bibliografía entera, dentro de controles de
contenido (w:sdt), y python-docx no los lee. En un proyecto real eso produjo 26
"referencias no citadas", "0 citas en antecedentes" y 17 "espacios antes de signo"
que eran los huecos de las citas invisibles. Los documentos se arman acá con el
mismo XML que escribe Mendeley.
"""
import copy

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from core.citas import extraer_citas, revisar_citas
from core.utils import Obs, bloques_en_orden, desenvolver_controles

REGLAS = {"citas": {"estilo": "APA7", "seccion_referencias": "Referencias"}}


def _sdt_en_linea(texto):
    return parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:tag w:val="MENDELEY_CITATION"/></w:sdtPr>'
        f'<w:sdtContent><w:r><w:t xml:space="preserve">{texto}</w:t></w:r></w:sdtContent></w:sdt>')


def _parrafo_con_cita(doc, antes, cita, despues):
    p = doc.add_paragraph(antes)
    p._p.append(_sdt_en_linea(cita))
    p.add_run(despues)
    return p


def _bibliografia_en_bloque(doc, referencias):
    """la bibliografía de Mendeley: un solo control con un párrafo por referencia"""
    sdt = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr><w:tag w:val="MENDELEY_BIBLIOGRAPHY"/>'
                    f'</w:sdtPr><w:sdtContent/></w:sdt>')
    contenido = sdt.find(qn("w:sdtContent"))
    for texto in referencias:
        contenido.append(copy.deepcopy(doc.add_paragraph(texto)._p))
    for p in list(doc.element.body.iter(qn("w:p")))[-len(referencias):]:
        doc.element.body.remove(p)
    doc.element.body.insert(len(doc.element.body) - 1, sdt)   # antes de sectPr


def _documento(tmp_path):
    doc = Document()
    _parrafo_con_cita(doc, "El conjunto reunió 671 cuencas ", "(Addor et al., 2017)", ".")
    _parrafo_con_cita(doc, "", "Demšar (2006)", " estableció el procedimiento.")
    doc.add_paragraph("Referencias")
    _bibliografia_en_bloque(doc, [
        "Addor, N., Newman, A. J., y Clark, M. P. (2017). The CAMELS data set. HESS, 21(10), 5293-5313.",
        "Demšar, J. (2006). Statistical comparisons of classifiers. JMLR, 7(1), 1-30.",
    ])
    ruta = tmp_path / "mendeley.docx"
    doc.save(str(ruta))
    return Document(str(ruta))


def test_sin_desenvolver_python_docx_no_ve_las_citas(tmp_path):
    """el problema de fondo: el texto de la cita no está en Paragraph.text"""
    doc = _documento(tmp_path)
    textos = [b.text for b in bloques_en_orden(doc)]
    assert textos[0] == "El conjunto reunió 671 cuencas ."
    assert not any("Addor, N." in t for t in textos), "la bibliografía en bloque tampoco se ve"


def test_desenvolver_deja_el_texto_completo(tmp_path):
    doc = _documento(tmp_path)
    assert desenvolver_controles(doc) == 3
    textos = [b.text for b in bloques_en_orden(doc)]
    assert textos[0] == "El conjunto reunió 671 cuencas (Addor et al., 2017)."
    assert textos[1] == "Demšar (2006) estableció el procedimiento."
    assert any(t.startswith("Addor, N.") for t in textos), "la bibliografía tiene que ser párrafos"
    assert not list(doc.element.body.iter(qn("w:sdt")))


def test_las_citas_de_mendeley_emparejan_con_sus_referencias(tmp_path):
    doc = _documento(tmp_path)
    desenvolver_controles(doc)
    bloques = list(bloques_en_orden(doc))
    i_ref = next(i for i, b in enumerate(bloques) if b.text == "Referencias")
    obs = Obs()
    revisar_citas(bloques, {"Referencias": (i_ref, len(bloques))}, REGLAS, obs)
    problemas = [x["observacion"] for x in obs.items if "citad" in x["observacion"] or "Cita sin" in x["observacion"]]
    assert problemas == []


def test_un_control_anidado_tambien_se_desenvuelve(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Según ")
    externo = _sdt_en_linea("")
    externo.find(qn("w:sdtContent")).append(_sdt_en_linea("Zhang (2003)"))
    p._p.append(externo)
    assert desenvolver_controles(doc) == 2
    assert doc.paragraphs[0].text == "Según Zhang (2003)"


def test_apellidos_con_letras_de_otros_idiomas_y_con_inicial():
    """'Demšar' se cortaba en la 'š'; 'Sánchez D.' no se reconocía por la inicial"""
    assert extraer_citas("Demšar (2006) estableció")[0][:2] == ("demsar", "2006")
    assert extraer_citas("Sánchez D. y Laqui V. (2009) lo hicieron")[0][:2] == ("sanchez d", "2009")
    assert extraer_citas("(Llanque Chayña, 2022)")[0][:2] == ("llanque chayna", "2022")


# --------------------------------------------------------------------------- ortografía técnica

import pytest

from core.ortografia import es_termino_tecnico


@pytest.mark.parametrize("palabra, texto, esperado", [
    ("v1.0", "Se emplea CAMELS-PE v1.0 de acceso abierto", True),
    ("ETS", "ETS no admite regresores", True),
    ("SeasonalNaive", "El modelo SeasonalNaive se usa", True),
    ("PISCO_HyM_GR2M", "el modelo PISCO_HyM_GR2M mensual", True),
    ("hoc", "seguida del post\xa0hoc de Nemenyi", True),        # con espacio de no separación
    ("priori", "declarados a priori en el plan", True),
    ("t", "aproximación por prueba t pareada", True),
    ("hidro", "la variabilidad hidro climática", False),        # este sí es un error
    ("ESTADISTICA", "ESTADISTICA DESCRIPTIVA", False),          # mayúsculas largas no son sigla
    ("t", "la letra t sola", False),
])
def test_terminos_tecnicos_que_el_corrector_no_marca(palabra, texto, esperado):
    assert es_termino_tecnico(palabra, texto, texto.index(palabra)) is esperado


def test_se_aplican_las_palabras_que_agrega_el_revisor(tmp_path, monkeypatch):
    """
    El botón "Palabras permitidas" edita la copia de la carpeta de datos; antes solo se
    leía la lista interna del programa y, en el .exe, lo agregado nunca se aplicaba.
    """
    from docx import Document as Doc
    from core import config
    from core.ortografia import revisar_ortografia
    (tmp_path / "permitidas.txt").write_text("statsforecast\n", encoding="utf-8")
    monkeypatch.setattr(config, "DATOS", str(tmp_path))
    doc = Doc()
    doc.add_paragraph("Se usa la biblioteca statsforecast para los modelos locales.")
    bloques = list(bloques_en_orden(doc))
    filas = []
    reglas = {"ortografia": {"motor": "diccionario", "diccionario": "dic/es_PE"}}
    revisar_ortografia(bloques, {}, reglas, Obs(), filas)
    assert not any(f["palabra"] == "statsforecast" for f in filas), filas


# --------------------------------------------------------------------------- formato de referencias

from core.citas import defectos_de_referencia


@pytest.mark.parametrize("referencia, defecto", [
    ("Llauca, H. (2021). PISCO_HyM_GR2M. Water 2021, Vol. 13, Page 1048, 13(8), 1048. https://doi.org/10.3390/w13081048",
     "Mendeley"),
    ("Nearing, G. (2024). Global prediction of floods. Nature 2024 627:8004, 627(8004), 559-563.", "exportación"),
    ("Reyes, T. (2019). Aplicación de modelos. Aporte Santiaguino, 12(1), ág: 70-80.", "Páginas mal escritas"),
    ("Béjar, W. (2016). Predicción de caudales y el modelo ANFIS Infere…. Research in Computing Science, 113.",
     "cortada"),
    ("Llanque, E. (2022). MODELACION HIDROLÓGICA CON PRECIPITACIONES OBTENIDAS POR SATÉLITE. Revista, 11(4), 214-226.",
     "mayúsculas"),
    ("Lujano, E. (2014). Pronóstico de caudales. Revista de Investigaciones Altoandinas, 16(1). https://doi.org/10.18271/ria.2014.93",
     "sin páginas"),
])
def test_defectos_de_forma_de_una_referencia(referencia, defecto):
    assert any(defecto in o for o, _ in defectos_de_referencia(referencia))


@pytest.mark.parametrize("referencia", [
    "Zhang, P. G. (2003). Time series forecasting using a hybrid ARIMA and neural network model. Neurocomputing, 50, 159-175. https://doi.org/10.1016/S0925-2312(01)00702-0",
    "Llauca, H. (2023). Construction of a daily streamflow dataset. Journal of Hydrology: Regional Studies, 47, 101381. https://doi.org/10.1016/J.EJRH.2023.101381",
    "Demšar, J. (2006). Statistical comparisons of classifiers. Journal of Machine Learning Research, 7(1), 1-30. http://jmlr.org/papers/v7/demsar06a.html",
    "Instituto Nacional de Estadística e Informática [INEI]. (2020). Perú: Encuesta Demográfica y de Salud Familiar ENDES 2019. INEI.",
    "Hyndman, R. J. (2021). Forecasting: principles and practice (3.a ed.). OTexts. https://otexts.com/fpp3/...",
])
def test_una_referencia_bien_formada_no_se_observa(referencia):
    assert defectos_de_referencia(referencia) == []
