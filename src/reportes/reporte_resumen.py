"""
Hoja de revisión resumida, en el formato que usa la Dirección de Investigación:
título, ficha del proyecto, asesor y una lista corta de observaciones,
cada una anclada a la primera línea donde aparece el problema.
"""
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TINTA = "1A1A1A"
ACENTO = "1F3864"
SUAVE = "6E6E6E"
LINEA = "BFBFBF"
FUENTE = "Arial"


def _fuente(run):
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), FUENTE)


def txt(par, texto, negrita=False, tam=11, color=TINTA, mayus=False):
    r = par.add_run(texto)
    r.bold = negrita
    r.font.size = Pt(tam)
    r.font.name = FUENTE
    r.font.color.rgb = RGBColor.from_string(color)
    if mayus:
        r.font.all_caps = True
    _fuente(r)
    return r


def sombrear(par, color):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    ppr.append(shd)


def borde_inferior(par, color=LINEA, grosor=6):
    ppr = par._p.get_or_add_pPr()
    bordes = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(grosor))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    bordes.append(b)
    ppr.append(bordes)


# --------------------------------------------------------------------------- redacción de observaciones

def _linea_de(items, filtro):
    """primer número de línea entre las observaciones que cumplen el filtro"""
    nums = sorted(it["nlinea"] for it in items if filtro(it) and it.get("nlinea"))
    return nums[0] if nums else None


def _cuantas(items, filtro):
    return sum(1 for it in items if filtro(it))


def _frase(numero, texto, varias=True):
    if numero:
        return f"Línea {numero}{' y otros' if varias else ''}: {texto}"
    return texto[0].upper() + texto[1:]


CITA_TEXTO = ("Cita sin referencia", "'et al", "Cita sin coma", "citas numéricas")
REFERENCIA = ("Referencia sin año", "sangría francesa", "orden alfabético", "DOI", "Referencia no citada",
              "no tiene entradas")


def redactar(items, filas_orto, extras=(), tipo="proyecto"):
    """convierte las observaciones detalladas en la lista corta de la hoja de revisión"""
    obs = []
    esquema = _cuantas(items, lambda x: x["categoria"] == "Estructura" and x["severidad"] == "Error")
    formato = [x for x in items if x["categoria"] == "Formato" and x["severidad"] == "Error"]
    guia = [x for x in items if "texto guía" in x["observacion"]]
    extension = [x for x in items if x["categoria"] == "Extensión" and x["severidad"] == "Error"]

    if esquema:
        obs.append(f"Revisar y cumplir el esquema de {tipo} de tesis PGI")
    if [x for x in formato if x not in guia]:
        detalle = ", ".join(sorted({_que_formato(x["observacion"]) for x in formato if x not in guia}))
        obs.append(f"Corregir el formato del documento: {detalle}")
    if guia:
        n = _linea_de(guia, lambda x: True)
        obs.append(_frase(n, "eliminar el texto guía de la plantilla que quedó sin borrar",
                          varias=len(guia) > 1))
    if filas_orto:
        obs.append("Revisión gramatical y ortográfica de todo el proyecto")

    n = _linea_de(items, lambda x: x["categoria"] == "Citas" and any(c in x["observacion"] for c in CITA_TEXTO))
    if n or _cuantas(items, lambda x: x["categoria"] == "Citas" and any(c in x["observacion"] for c in CITA_TEXTO)):
        obs.append(_frase(n, "corregir la forma de citar, estilo APA 7ª ed"))

    if filas_orto:
        nums = sorted(f["nlinea"] for f in filas_orto if f.get("nlinea"))
        obs.append(_frase(nums[0] if nums else None, "corregir ortografía y gramática",
                          varias=len(filas_orto) > 1))

    n = _linea_de(items, lambda x: x["categoria"] == "Citas" and any(c in x["observacion"] for c in REFERENCIA))
    if n or _cuantas(items, lambda x: x["categoria"] == "Citas" and any(c in x["observacion"] for c in REFERENCIA)):
        obs.append(_frase(n, "referenciar según norma APA 7ª ed"))

    for x in extension:
        obs.append(_redactar_extension(x["observacion"]))

    aviso = [x for x in items if "No se pudo verificar las citas" in x["observacion"]]
    if aviso:
        obs.append("Titular la sección de referencias como indica el esquema; no se pudo verificar "
                   "la correspondencia entre citas y referencias")

    for x in [i for i in items if i["categoria"] == "Revisor"]:
        obs.append(_frase(x.get("nlinea"), x["observacion"][0].lower() + x["observacion"][1:])
                   if x.get("nlinea") else x["observacion"])

    obs += [e.strip() for e in extras if e.strip()]
    return obs


def _redactar_extension(o):
    m = re.search(r"El documento tiene (\d+) páginas, máximo (\d+)", o)
    if m:
        return f"Reducir la extensión a {m.group(2)} páginas como máximo (tiene {m.group(1)})"
    m = re.search(r"El título tiene (\d+) palabras, máximo (\d+)", o)
    if m:
        return f"Reducir el título a {m.group(2)} palabras como máximo (tiene {m.group(1)})"
    m = re.search(r"Tiene (\d+) palabras clave, máximo (\d+)", o)
    if m:
        return f"Dejar {m.group(2)} palabras clave como máximo (tiene {m.group(1)})"
    return o


def _que_formato(observacion):
    o = observacion.lower()
    if "margen" in o:
        return "márgenes"
    if "fuente" in o:
        return "tipo de letra"
    if "tamaño" in o:
        return "tamaño de letra"
    if "interlineado" in o:
        return "interlineado"
    if "justificar" in o:
        return "texto justificado"
    if "encabezado" in o:
        return "encabezado institucional"
    if "papel" in o:
        return "tamaño de papel"
    return "formato general"


# --------------------------------------------------------------------------- documento

def escribir_resumen(ruta_salida, datos, observaciones):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    cab = f"REVISIÓN PGI {datos.get('fecha_revision', '')}"
    if datos.get("n_revision"):
        cab += f" – {datos['n_revision']}"
    p = doc.add_paragraph()
    txt(p, cab, negrita=True, tam=12)

    # ficha oscura del proyecto
    ficha = doc.add_paragraph()
    ficha.paragraph_format.space_before = Pt(6)
    ficha.paragraph_format.space_after = Pt(0)
    ficha.paragraph_format.left_indent = Cm(0.2)
    sombrear(ficha, "1A1A1A")
    if datos.get("codigo"):
        txt(ficha, datos["codigo"] + "\n", negrita=True, tam=11, color="FFFFFF")
    txt(ficha, (datos.get("proyecto") or "(título del proyecto)").upper(), tam=8, color="D9D9D9")
    pie = doc.add_paragraph()
    pie.paragraph_format.space_after = Pt(14)
    pie.paragraph_format.left_indent = Cm(0.2)
    sombrear(pie, "1A1A1A")
    etapa = datos.get("etapa") or "Etapa 2: Revisión de formato de proyecto"
    txt(pie, f"{etapa}   ·   1 tesista   ·   {datos.get('archivo', '')}", tam=8, color="BFBFBF")

    if datos.get("tesista"):
        p = doc.add_paragraph()
        txt(p, "Tesista: ", negrita=True)
        txt(p, datos["tesista"])
    p = doc.add_paragraph()
    txt(p, "Asesor: ", negrita=True)
    txt(p, datos.get("asesor") or "(completar)", color=TINTA if datos.get("asesor") else SUAVE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    txt(p, "OBSERVACIONES", negrita=True)
    borde_inferior(p)

    for o in observaciones:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(6)
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        m = re.match(r"^(Línea \d+(?: y otros)?):\s*(.+)$", o)
        if m:
            txt(q, m.group(1) + ": ", negrita=True)
            txt(q, m.group(2))
        else:
            txt(q, o)
    if not observaciones:
        q = doc.add_paragraph()
        txt(q, "Sin observaciones de formato. Pasa a la siguiente etapa.", color=SUAVE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    txt(p, "_" * 42, tam=10, color=LINEA)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    txt(p, datos.get("revisor") or "Revisor", negrita=True, tam=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    txt(p, "Dirección de Investigación – FINESI", tam=9, color=SUAVE)

    doc.save(ruta_salida)
