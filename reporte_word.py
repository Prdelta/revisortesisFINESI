"""
Documento Word de observaciones. Diseño formal, minimalista y funcional:
una sola tipografía, una tinta de acento, reglas finas en lugar de bloques de color.
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TINTA = "1A1A1A"        # texto
ACENTO = "1F3864"       # azul institucional, sobrio
SUAVE = "6E6E6E"        # texto secundario
LINEA = "BFBFBF"        # bordes
FONDO = "F2F2F2"        # relleno de encabezados de tabla
SEVERIDAD = {"Error": "8B1A1A", "Advertencia": "8A6D00", "Revisar": "3C5A78"}
FUENTE = "Arial"


# --------------------------------------------------------------------------- utilidades de formato

def _rfonts(run):
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), FUENTE)


def txt(par, texto, negrita=False, tam=9.5, color=TINTA, espaciado=None, cursiva=False):
    r = par.add_run(texto)
    r.bold = negrita
    r.italic = cursiva
    r.font.size = Pt(tam)
    r.font.name = FUENTE
    r.font.color.rgb = RGBColor.from_string(color)
    _rfonts(r)
    if espaciado:  # separación entre letras, en puntos
        rpr = r._r.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(espaciado * 20)))
        rpr.append(sp)
    return r


def parrafo(doc, texto="", negrita=False, tam=9.5, color=TINTA, alineacion=None,
            antes=0, despues=2, espaciado=None, cursiva=False):
    p = doc.add_paragraph()
    if alineacion is not None:
        p.alignment = alineacion
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    if texto:
        txt(p, texto, negrita, tam, color, espaciado, cursiva)
    return p


def regla(par, grosor=6, color=ACENTO, arriba=False):
    """línea horizontal fina como borde del párrafo"""
    pPr = par._p.get_or_add_pPr()
    bordes = pPr.find(qn("w:pBdr"))
    if bordes is None:
        bordes = OxmlElement("w:pBdr")
        pPr.append(bordes)
    b = OxmlElement("w:top" if arriba else "w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(grosor))
    b.set(qn("w:space"), "2")
    b.set(qn("w:color"), color)
    bordes.append(b)


def sombrear(celda_, color):
    tcPr = celda_._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def bordes_tabla(tabla, color=LINEA, grosor=4, solo_horizontales=False):
    tblPr = tabla._tbl.tblPr
    viejo = tblPr.find(qn("w:tblBorders"))
    if viejo is not None:
        tblPr.remove(viejo)
    bordes = OxmlElement("w:tblBorders")
    activos = ["top", "bottom", "insideH"] if solo_horizontales else \
              ["top", "left", "bottom", "right", "insideH", "insideV"]
    for lado in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        b = OxmlElement("w:" + lado)
        if lado in activos:
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), str(grosor))
            b.set(qn("w:color"), color)
        else:
            b.set(qn("w:val"), "none")
            b.set(qn("w:sz"), "0")
        b.set(qn("w:space"), "0")
        bordes.append(b)
    tblPr.append(bordes)


def anchos_fijos(tabla, anchos):
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tabla._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, ancho in zip(grid.findall(qn("w:gridCol")), anchos):
            col.set(qn("w:w"), str(int(ancho.cm * 567)))
    for fila in tabla.rows:
        for c, ancho in zip(fila.cells, anchos):
            c.width = ancho


def margen_celdas(tabla, arriba=40, abajo=40, lados=80):
    tblPr = tabla._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for lado, valor in (("top", arriba), ("bottom", abajo), ("left", lados), ("right", lados)):
        e = OxmlElement("w:" + lado)
        e.set(qn("w:w"), str(valor))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def repetir_encabezado(fila):
    trPr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def celda(c, texto, negrita=False, tam=8.5, color=TINTA, alineacion=None):
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if alineacion is not None:
        p.alignment = alineacion
    txt(p, texto, negrita, tam, color)
    return p


def pie_de_pagina(seccion, texto_izquierda):
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    regla(p, 4, LINEA, arriba=True)
    txt(p, texto_izquierda, False, 7.5, SUAVE)
    r = p.add_run()
    _rfonts(r)
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(SUAVE)
    tab = OxmlElement("w:ptab")
    tab.set(qn("w:relativeTo"), "margin")
    tab.set(qn("w:alignment"), "right")
    tab.set(qn("w:leader"), "none")
    r._r.append(tab)
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    r._r.append(inicio)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r._r.append(instr)
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    r._r.append(fin)


# --------------------------------------------------------------------------- documento

def escribir_word(ruta_salida, datos, obs_items, filas_orto, motor, resumen_categorias, mapa=None):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = Cm(1.9)
    s.bottom_margin = Cm(1.7)
    s.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(TINTA)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    pie_de_pagina(s, f"Dirección de Investigación · FINESI · {datos.get('archivo', '')}")

    # ---- membrete
    parrafo(doc, "UNIVERSIDAD NACIONAL DEL ALTIPLANO – PUNO", True, 10, TINTA,
            WD_ALIGN_PARAGRAPH.CENTER, espaciado=0.6, despues=0)
    parrafo(doc, "Facultad de Ingeniería Estadística e Informática", False, 9, SUAVE,
            WD_ALIGN_PARAGRAPH.CENTER, despues=0)
    p = parrafo(doc, "Dirección de Investigación", False, 9, SUAVE, WD_ALIGN_PARAGRAPH.CENTER, despues=6)
    regla(p, 8, ACENTO)

    parrafo(doc, f"HOJA DE OBSERVACIONES · {datos.get('tipo', '').upper()} DE TESIS", True, 12.5,
            ACENTO, WD_ALIGN_PARAGRAPH.CENTER, antes=8, despues=10, espaciado=0.8)

    # ---- datos del expediente
    campos = [("Tesista", datos.get("tesista")),
              ("Título", datos.get("proyecto")),
              ("Asesor", datos.get("asesor")),
              ("Presentación", datos.get("fecha_subida")),
              ("Archivo", datos.get("archivo")),
              ("Revisión", datos.get("fecha_revision"))]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for etiqueta, valor in campos:
        fila = t.add_row()
        celda(fila.cells[0], etiqueta.upper(), True, 8, SUAVE)
        celda(fila.cells[1], valor or "(completar)", False, 9, TINTA if valor else SUAVE)
    bordes_tabla(t, LINEA, 4, solo_horizontales=True)
    margen_celdas(t, 55, 55, 0)
    anchos_fijos(t, [Cm(3.2), Cm(13.4)])

    # ---- resumen
    n = {k: sum(1 for x in obs_items if x["severidad"] == k) for k in ("Error", "Advertencia", "Revisar")}
    p = parrafo(doc, "", antes=12, despues=1)
    txt(p, f"{n['Error']} ", True, 11, SEVERIDAD["Error"])
    txt(p, "errores     ", False, 9, SUAVE)
    txt(p, f"{n['Advertencia']} ", True, 11, SEVERIDAD["Advertencia"])
    txt(p, "advertencias     ", False, 9, SUAVE)
    txt(p, f"{n['Revisar']} ", True, 11, SEVERIDAD["Revisar"])
    txt(p, "por verificar", False, 9, SUAVE)
    detalle = "     ·     ".join(f"{c}: {v}" for c, v in resumen_categorias.items() if v)
    if detalle:
        parrafo(doc, detalle, False, 8, SUAVE, despues=6)

    parrafo(doc, "Observaciones generadas automáticamente sobre formato, estructura, citas y ortografía. "
                 "Requieren validación del revisor antes de ser comunicadas al tesista.",
            False, 8, SUAVE, cursiva=True, despues=10)

    # ---- observaciones
    p = parrafo(doc, "OBSERVACIONES", True, 9.5, ACENTO, antes=6, despues=4, espaciado=1.0)
    regla(p, 6, ACENTO)
    if mapa is not None and not mapa:
        parrafo(doc, "No se pudo determinar página y línea (LibreOffice no está instalado). "
                     "Se indica la sección y el texto de referencia.", False, 8, SUAVE, cursiva=True)

    anchos = [Cm(0.8), Cm(1.9), Cm(2.2), Cm(3.7), Cm(4.0), Cm(4.0)]
    t = doc.add_table(rows=1, cols=6)
    for c, texto in zip(t.rows[0].cells, ["N°", "TIPO", "PÁG./LÍNEA", "UBICACIÓN", "OBSERVACIÓN", "CORRECCIÓN"]):
        sombrear(c, FONDO)
        celda(c, texto, True, 7.5, ACENTO)
    repetir_encabezado(t.rows[0])
    for k, it in enumerate(obs_items, 1):
        fila = t.add_row()
        valores = [str(k), it["severidad"], it.get("linea", "") or "—",
                   it["ubicacion"], it["observacion"], it["detalle"] or ""]
        for i, (c, valor) in enumerate(zip(fila.cells, valores)):
            if i == 1:
                celda(c, valor, True, 8, SEVERIDAD.get(it["severidad"], TINTA))
            elif i in (0, 3):
                celda(c, valor, False, 8, SUAVE)
            else:
                celda(c, valor, False, 8.5)
    if not obs_items:
        celda(t.add_row().cells[0], "Sin observaciones automáticas.", False, 9)
    bordes_tabla(t, LINEA, 4, solo_horizontales=True)
    margen_celdas(t, 50, 50, 70)
    anchos_fijos(t, anchos)

    # ---- ortografía
    if filas_orto:
        p = parrafo(doc, "ORTOGRAFÍA Y GRAMÁTICA", True, 9.5, ACENTO, antes=14, despues=4, espaciado=1.0)
        regla(p, 6, ACENTO)
        parrafo(doc, f"Motor: {motor}. Puede señalar términos técnicos correctos; verificar antes de observar.",
                False, 8, SUAVE, cursiva=True, despues=4)
        anchos_o = [Cm(2.2), Cm(3.4), Cm(2.6), Cm(4.7), Cm(3.7)]
        t = doc.add_table(rows=1, cols=5)
        for c, texto in zip(t.rows[0].cells, ["PÁG./LÍNEA", "UBICACIÓN", "TEXTO", "PROBLEMA", "SUGERENCIA"]):
            sombrear(c, FONDO)
            celda(c, texto, True, 7.5, ACENTO)
        repetir_encabezado(t.rows[0])
        for f in filas_orto:
            fila = t.add_row()
            celda(fila.cells[0], f.get("linea", "") or "—", False, 8)
            celda(fila.cells[1], str(f["parrafo"]), False, 8, SUAVE)
            celda(fila.cells[2], str(f["palabra"]), True, 8.5)
            celda(fila.cells[3], str(f["problema"]), False, 8)
            celda(fila.cells[4], str(f["sugerencias"]), False, 8.5)
        bordes_tabla(t, LINEA, 4, solo_horizontales=True)
        margen_celdas(t, 50, 50, 70)
        anchos_fijos(t, anchos_o)

    # ---- cierre
    p = parrafo(doc, "OBSERVACIONES DEL REVISOR", True, 9.5, ACENTO, antes=16, despues=4, espaciado=1.0)
    regla(p, 6, ACENTO)
    for _ in range(4):
        q = parrafo(doc, "", despues=9)
        regla(q, 4, LINEA)

    p = parrafo(doc, "", antes=30, despues=0, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.left_indent = Cm(5.5)
    p.paragraph_format.right_indent = Cm(5.5)
    regla(p, 4, TINTA)
    parrafo(doc, datos.get("revisor") or "Revisor responsable", True, 9, TINTA,
            WD_ALIGN_PARAGRAPH.CENTER, despues=0)
    parrafo(doc, "Dirección de Investigación – FINESI", False, 8, SUAVE, WD_ALIGN_PARAGRAPH.CENTER)

    carpeta = os.path.dirname(os.path.abspath(ruta_salida))
    os.makedirs(carpeta, exist_ok=True)
    doc.save(ruta_salida)
