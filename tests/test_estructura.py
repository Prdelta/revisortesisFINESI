"""
Títulos de sección: subtítulos internos que se contaban como la sección repetida, y
títulos cortados que se aceptaban sin aviso.
"""
from docx import Document

from core.estructura import detectar_secciones, revisar_estructura
from core.orquestador import cargar_reglas
from core.utils import Obs, bloques_en_orden


def _revisar(titulos_y_textos):
    doc = Document()
    for texto in titulos_y_textos:
        doc.add_paragraph(texto)
    bloques = list(bloques_en_orden(doc))
    obs = Obs()
    reglas = cargar_reglas("proyecto")
    revisar_estructura(bloques, reglas, obs, detectar_secciones(bloques, reglas))
    return [x["observacion"] for x in obs.items]


def test_los_subtitulos_de_hipotesis_no_son_la_seccion_repetida():
    obs = _revisar(["Hipótesis del trabajo", "HIPÓTESIS GENERAL", "El paradigma global supera al local.",
                    "HIPÓTESIS ESPECÍFICAS", "H1. El paradigma global supera al local."])
    assert not any("aparece" in o for o in obs), obs


def test_un_titulo_alargado_sigue_reconociendose_si_es_el_primero():
    """la regla del título alargado no se pierde: 'Referencias bibliográficas y fuentes'"""
    doc = Document()
    doc.add_paragraph("Referencias bibliográficas consultadas")
    reglas = cargar_reglas("proyecto")
    unicos, _, _ = detectar_secciones(list(bloques_en_orden(doc)), reglas)
    assert [n for _, n in unicos] == ["Referencias"]


def test_un_titulo_cortado_se_avisa():
    obs = _revisar(["Uso de los resultados y contribuciones del"])
    assert any("incompleto" in o for o in obs), obs


def test_el_titulo_exacto_y_el_alias_no_se_avisan():
    obs = _revisar(["Uso de los resultados y contribuciones del proyecto", "Recursos"])
    assert not any("incompleto" in o or "distinto" in o for o in obs), obs


def test_un_titulo_alargado_no_se_avisa():
    obs = _revisar(["Justificación del proyecto de investigación"])
    assert not any("incompleto" in o or "distinto" in o for o in obs), obs


def test_un_proyecto_con_anexos_no_parece_un_borrador():
    """
    La regla propia del revisor pide la matriz de consistencia en anexos, y a la vez
    'anexos' contaba como señal de borrador: un proyecto con Anexos y un subtítulo
    'Resultados esperados' quedaba sin revisar.
    """
    from core.orquestador import parece_otro_tipo
    doc = Document()
    for t in ("Anexos", "Matriz de consistencia", "Resultados"):
        doc.add_paragraph(t)
    ajenos = parece_otro_tipo(list(bloques_en_orden(doc)), cargar_reglas("proyecto"))
    assert "Anexos" not in ajenos and len(ajenos) < 2
