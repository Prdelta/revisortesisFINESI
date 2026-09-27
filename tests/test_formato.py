"""
Tamaño de letra en tablas y notas. APA admite un tamaño menor ahí; antes toda tabla
en 9 pt salía como Error de formato aunque el cuerpo estuviera bien.
"""
from docx import Document
from docx.shared import Pt

from core.formato import es_nota_de_tabla, revisar_formato
from core.orquestador import cargar_reglas
from core.utils import Obs, bloques_en_orden


def _documento(tam_cuerpo, tam_tabla, tam_nota):
    doc = Document()
    for estilo in (doc.styles["Normal"],):
        estilo.font.name, estilo.font.size = "Arial", Pt(10)
    doc.add_paragraph().add_run("Texto del cuerpo del proyecto. " * 20).font.size = Pt(tam_cuerpo)
    tabla = doc.add_table(rows=2, cols=2)
    for fila in tabla.rows:
        for celda in fila.cells:
            celda.paragraphs[0].add_run("Dato de la tabla con varias palabras").font.size = Pt(tam_tabla)
    doc.add_paragraph().add_run("Nota. Elaboración propia a partir de los datos. " * 3).font.size = Pt(tam_nota)
    return doc


def _tamanos(doc, tamano_tablas=9):
    reglas = cargar_reglas("proyecto")
    reglas["formato"] = dict(reglas["formato"], tamano_tablas_pt=tamano_tablas)
    obs = Obs()
    revisar_formato(doc, list(bloques_en_orden(doc)), {}, reglas, obs, ([], [], []))
    return [x["observacion"] for x in obs.items if "tamaño" in x["observacion"]]


def test_tabla_y_nota_en_9_pt_se_aceptan():
    assert _tamanos(_documento(10, 9, 9)) == []


def test_sin_la_clave_la_tabla_en_9_pt_se_observa():
    assert any("9 pt" in o for o in _tamanos(_documento(10, 9, 9), tamano_tablas=None))


def test_el_cuerpo_en_9_pt_se_sigue_observando():
    assert any("9 pt" in o for o in _tamanos(_documento(9, 10, 10)))


def test_una_tabla_en_8_pt_se_observa():
    assert any("8 pt" in o for o in _tamanos(_documento(10, 8, 10)))


def test_que_es_una_nota_de_tabla():
    assert es_nota_de_tabla("Nota. AA = aprendizaje automático")
    assert es_nota_de_tabla("Fuente: INEI (2020)")
    assert not es_nota_de_tabla("La nota promedio de los estudiantes fue 14")
