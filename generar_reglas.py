"""
Genera un archivo de reglas a partir de una plantilla Word (.docx).

Uso:
    python generar_reglas.py plantillas/borrador.docx --tipo borrador

Crea reglas/borrador.yaml con: esquema de secciones y subsecciones, margenes, fuente,
tamano, interlineado y encabezado, todo leido de la plantilla.
Es un PUNTO DE PARTIDA: hay que revisarlo contra el reglamento antes de usarlo.
"""
import argparse
import os
import re
import sys
from collections import Counter

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from revisor import DATOS as BASE, preparar_datos  # noqa: E402
from revisor import Estilos, limpiar_titulo  # noqa: E402

RE_CAPITULO = re.compile(r"^\s*CAP[IÍ]TULO\s+([IVXLC]+|\d+)\b", re.I)
RE_NUMERADO = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+\S")


def texto_titulo(txt):
    """quita numeracion y el texto guia entre parentesis, conserva mayusculas/tildes originales"""
    txt = re.sub(r"^\s*(?:[IVXLC]+|\d+(?:\.\d+)*|[a-z])[\.\)\-:]?\s+", "", txt.strip())
    txt = re.split(r"\(", txt, maxsplit=1)[0]
    return re.sub(r"\s+", " ", txt).strip(" .:")


def nivel_de(p):
    """nivel de titulo (0 = seccion principal, 1+ = subseccion) o None si no parece titulo"""
    txt = p.text.strip()
    if not txt or len(texto_titulo(txt)) > 90:
        return None
    nombre = (p.style.name or "").lower()
    m = re.search(r"(heading|t[ií]tulo)\s*(\d)", nombre)
    if m:
        return int(m.group(2)) - 1
    if RE_CAPITULO.match(txt):
        return 0
    ilvl = p._p.xpath("./w:pPr/w:numPr/w:ilvl/@w:val")
    if ilvl:
        return int(ilvl[0])
    m = RE_NUMERADO.match(txt)
    negrita = any(r.bold for r in p.runs if r.text.strip())
    if m and negrita:
        return m.group(1).count(".")
    return None


def extraer(doc):
    candidatos = [(nivel_de(p), p) for p in doc.paragraphs]
    candidatos = [(n, p) for n, p in candidatos if n is not None]
    if not candidatos:
        return [], {}
    minimo = min(n for n, _ in candidatos)
    secciones, subs = [], {}
    actual = None
    ultimo_fue_capitulo = False
    parrafos = doc.paragraphs
    for n, p in candidatos:
        nombre = texto_titulo(p.text)
        if RE_CAPITULO.match(p.text):
            # "CAPITULO I" suele ir seguido del nombre del capitulo en el siguiente parrafo
            idx = next(k for k, q in enumerate(parrafos) if q._p is p._p)
            sig = next((q for q in parrafos[idx + 1:idx + 4] if q.text.strip()), None)
            resto = RE_CAPITULO.sub("", p.text).strip(" :.-")
            nombre = f"{p.text.strip()}" if not resto and not sig else (resto or texto_titulo(sig.text))
        if not nombre:
            continue
        if n == minimo:
            actual = nombre
            if actual not in secciones:
                secciones.append(actual)
        elif actual:
            subs.setdefault(actual, [])
            if nombre not in subs[actual]:
                subs[actual].append(nombre)
    return secciones, subs


def formato(doc):
    s = doc.sections[0]
    est = Estilos(doc)
    fuentes, tamanos, inter = Counter(), Counter(), Counter()
    just = Counter()
    for p in doc.paragraphs:
        for r in p.runs:
            n = len(r.text.strip())
            if n:
                fuentes[est.fuente(r, p) or "?"] += n
                tamanos[est.tamano(r, p)] += n
        if len(p.text.strip()) > 60:
            inter[est.interlineado(p)] += 1
            just[est.alineacion(p) == WD_ALIGN_PARAGRAPH.JUSTIFY] += 1
    enc = [t.strip() for t in re.split(r"[\n\r]+", " \n".join(q.text for q in s.header.paragraphs)) if len(t.strip()) > 5]
    return dict(
        ancho=round(s.page_width.cm, 1), alto=round(s.page_height.cm, 1),
        izq=round(s.left_margin.cm, 2), der=round(s.right_margin.cm, 2),
        sup=round(s.top_margin.cm, 2), inf=round(s.bottom_margin.cm, 2),
        fuente=fuentes.most_common(1)[0][0] if fuentes else "Arial",
        tamano=tamanos.most_common(1)[0][0] if tamanos else 12,
        interlineado=inter.most_common(1)[0][0] if inter else 1.0,
        justificado=(just[True] >= just[False]) if just else True,
        encabezado=enc,
    )


# variantes frecuentes con las que los tesistas titulan cada sección
ALIAS_FRECUENTES = {
    "referencia": ["referencias", "referencias bibliograficas", "bibliografia", "fuentes consultadas",
                   "fuentes de informacion"],
    "bibliograf": ["referencias", "referencias bibliograficas", "bibliografia"],
    "metodolog": ["metodologia", "metodologia de la investigacion", "materiales y metodos",
                  "metodos y materiales"],
    "antecedente": ["antecedentes", "antecedentes de la investigacion", "estado del arte"],
    "justificac": ["justificacion", "justificacion de la investigacion", "justificacion del estudio"],
    "hipotesis": ["hipotesis", "hipotesis de la investigacion", "hipotesis de trabajo"],
    "objetivo general": ["objetivo general", "objetivo principal"],
    "objetivos especificos": ["objetivos especificos", "objetivo especifico"],
    "palabras": ["palabras clave", "palabras claves", "keywords"],
    "resumen": ["resumen", "sumario"],
    "abstract": ["abstract", "summary"],
    "conclusion": ["conclusiones", "conclusion"],
    "recomendacion": ["recomendaciones", "recomendacion"],
    "resultado": ["resultados", "resultados y discusion"],
    "discusion": ["discusion", "resultados y discusion"],
    "marco teorico": ["marco teorico", "bases teoricas", "revision de literatura"],
    "impacto": ["impactos esperados", "resultados esperados"],
    "presupuesto": ["presupuesto", "presupuesto del proyecto", "financiamiento"],
    "cronograma": ["cronograma", "cronograma de actividades"],
    "anexo": ["anexos", "apendices"],
    "introduccion": ["introduccion"],
}


def alias_de(nombre):
    """propone alias a partir de variantes frecuentes; el revisor los completa a mano"""
    clave = _norm_simple(nombre)
    propuestos = []
    for gatillo, variantes in ALIAS_FRECUENTES.items():
        if gatillo in clave:
            propuestos += [v for v in variantes if v != clave]
    vistos = []
    for v in propuestos:
        if v not in vistos:
            vistos.append(v)
    return vistos[:6]


def _norm_simple(t):
    import unicodedata
    t = unicodedata.normalize("NFD", t.lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9ñ ]+", " ", t)).strip()


def q(x):
    return '"' + str(x).replace('"', "'") + '"'


def escribir_yaml(tipo, ruta_plantilla, secciones, subs, f):
    ref = next((x for x in secciones if re.search(r"referencia|bibliograf", x, re.I)), None)
    otro = ["presupuesto", "cronograma de actividades", "cronograma"] if "borrador" in tipo.lower() else \
           ["resultados", "discusion", "conclusiones", "recomendaciones", "resultados y discusion"]
    lin = []
    lin.append(f"# Reglas de revision: {tipo.upper()} (generado automaticamente desde {os.path.basename(ruta_plantilla)})")
    lin.append("# REVISAR TODO contra el reglamento antes de usar. Los valores salen de como esta hecha la plantilla,")
    lin.append("# no de una norma escrita. Borra o corrige lo que no corresponda.")
    lin.append(f"tipo: {tipo}")
    lin.append("")
    lin.append(f"plantilla: {os.path.relpath(ruta_plantilla, BASE).replace(os.sep, '/')}")
    lin.append("")
    lin.append("formato:")
    lin.append(f"  papel: {{ancho_cm: {f['ancho']}, alto_cm: {f['alto']}}}")
    lin.append(f"  margenes_cm: {{izquierdo: {f['izq']}, derecho: {f['der']}, superior: {f['sup']}, inferior: {f['inf']}}}")
    lin.append("  tolerancia_margen_cm: 0.1")
    lin.append(f"  fuente: {f['fuente']}")
    lin.append(f"  tamano_pt: {f['tamano']:g}")
    lin.append(f"  interlineado: {f['interlineado']}")
    lin.append(f"  alineacion_cuerpo: {'justificado' if f['justificado'] else 'izquierda'}")
    lin.append("  tolerancia_fuera_de_norma: 0.03")
    lin.append("  encabezado_debe_contener:" + ("" if f["encabezado"] else " []"))
    for e in f["encabezado"]:
        lin.append(f"    - {q(e)}")
    lin.append("")
    lin.append("extension:")
    lin.append("  max_paginas: null            # COMPLETAR si el reglamento pone limite (null = no se revisa)")
    lin.append("  seccion_titulo: null         # nombre de la seccion que contiene el titulo, si hay limite de palabras")
    lin.append("  titulo_max_palabras: null")
    lin.append(f"  seccion_palabras_clave: {q(next((x for x in secciones if re.search('clave', x, re.I)), '')) if any(re.search('clave', x, re.I) for x in secciones) else 'null'}")
    lin.append("  palabras_clave_max: null")
    lin.append("")
    lin.append("# Esquema en orden. En 'alias' van otras formas validas de titular la seccion.")
    lin.append("# Los alias vienen propuestos; agrega los que veas en la practica y borra los que no apliquen.")
    lin.append("secciones:")
    for s in secciones:
        al = ", ".join(q(x) for x in alias_de(s))
        lin.append(f"  - {{nombre: {q(s)}, alias: [{al}]}}")
    if not secciones:
        lin.append("  []   # NO se detectaron titulos: escribe el esquema a mano")
    lin.append("")
    lin.append("subsecciones:" + ("" if subs else " {}"))
    for padre, hijos in subs.items():
        lin.append(f"  {q(padre)}:")
        for h in hijos:
            lin.append(f"    - {q(h)}")
    lin.append("")
    lin.append("secciones_con_tabla: []        # COMPLETAR, ej. [\"Cronograma de actividades\"]")
    lin.append("")
    lin.append(f"# Si aparecen estos titulos, el documento probablemente es de otro tipo y no se revisa")
    lin.append(f"secciones_de_otro_tipo: [{', '.join(q(x) for x in otro)}]")
    lin.append("")
    lin.append("citas:")
    lin.append("  estilo: APA7                # CONFIRMAR con la direccion")
    lin.append(f"  seccion_referencias: {q(ref) if ref else q('Referencias') + '   # REVISAR: no se encontro en la plantilla'}")
    lin.append("")
    lin.append("ortografia:")
    lin.append("  motor: auto")
    lin.append("  palabras_permitidas: permitidas.txt")
    lin.append("  diccionario: dic/es_PE")
    lin.append(f"  omitir_secciones: [{q(ref) if ref else ''}]")
    return "\n".join(lin) + "\n"


def main():
    ap = argparse.ArgumentParser(description="Genera reglas de revision desde una plantilla .docx")
    ap.add_argument("plantilla")
    ap.add_argument("--tipo", required=True, help="nombre del tipo, ej. borrador")
    ap.add_argument("--sobrescribir", action="store_true")
    a = ap.parse_args()
    if not os.path.exists(a.plantilla):
        sys.exit(f"No existe: {a.plantilla}")
    preparar_datos()
    destino = os.path.join(BASE, "reglas", f"{a.tipo}.yaml")
    if os.path.exists(destino) and not a.sobrescribir:
        with open(destino, encoding="utf8") as fh:
            if "secciones:" in fh.read().replace("secciones: []", ""):
                sys.exit(f"{destino} ya existe y tiene reglas. Usa --sobrescribir si de verdad quieres reemplazarlo.")
    doc = Document(a.plantilla)
    secciones, subs = extraer(doc)
    f = formato(doc)
    ruta_plantilla = os.path.abspath(a.plantilla)
    with open(destino, "w", encoding="utf8") as fh:
        fh.write(escribir_yaml(a.tipo, ruta_plantilla, secciones, subs, f))
    print(f"Reglas generadas en {destino}")
    print(f"Secciones detectadas ({len(secciones)}):")
    for s in secciones:
        print(f"  - {s}" + (f"  ({len(subs[s])} subsecciones)" if s in subs else ""))
    print(f"Formato leído: {f['fuente']} {f['tamano']:g} pt, interlineado {f['interlineado']}, "
          f"márgenes {f['izq']}/{f['der']}/{f['sup']}/{f['inf']} cm")
    print("Abre el archivo, compáralo con el reglamento y completa lo marcado como COMPLETAR o REVISAR.")


if __name__ == "__main__":
    main()
