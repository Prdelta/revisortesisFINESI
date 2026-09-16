import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rapidfuzz import fuzz
import localizador
import reglas_revisor
from reporte_resumen import escribir_resumen, redactar
from reporte_word import escribir_word
from .utils import *

def revisar_formato(doc, bloques, rangos, reglas, obs):
    f = reglas.get("formato") or {}
    if not f:
        return
    est = Estilos(doc)

    for k, s in enumerate(doc.sections):
        lugar = f"Sección de página {k+1}"
        if abs(s.page_width.cm - f["papel"]["ancho_cm"]) > 0.2 or abs(s.page_height.cm - f["papel"]["alto_cm"]) > 0.2:
            obs.add("Formato", "Error", lugar, "Tamaño de papel incorrecto",
                    f"Tiene {s.page_width.cm:.1f} x {s.page_height.cm:.1f} cm, debe ser A4 (21 x 29.7 cm)")
        tol = f["tolerancia_margen_cm"]
        for nombre, valor in [("izquierdo", s.left_margin), ("derecho", s.right_margin),
                              ("superior", s.top_margin), ("inferior", s.bottom_margin)]:
            esperado = f["margenes_cm"][nombre]
            if valor is None or abs(valor.cm - esperado) > tol:
                real = f"{valor.cm:.2f}" if valor is not None else "?"
                obs.add("Formato", "Error", lugar, f"Margen {nombre} incorrecto",
                        f"Tiene {real} cm, debe ser {esperado} cm")
        enc = " ".join(p.text for p in s.header.paragraphs)
        for frase in f.get("encabezado_debe_contener", []):
            if norm(frase) not in norm(enc):
                obs.add("Formato", "Error", lugar, "Encabezado incompleto o modificado", f"Debe contener: '{frase}'")

    # fuente y tamano: conteo por caracteres
    fuentes, tamanos = Counter(), Counter()
    ejemplos_f, ejemplos_t = defaultdict(list), defaultdict(list)
    total = 0
    parrafos = []
    for i, b in enumerate(bloques):
        if isinstance(b, Paragraph):
            parrafos.append((i, b))
        else:
            for fila in b.rows:
                for c in fila.cells:
                    for p in c.paragraphs:
                        parrafos.append((i, p))
    titulos = {i for i, _ in detectar_secciones(bloques, reglas)[1]}
    for i, p in parrafos:
        es_titulo = i in titulos or re.search(r"(heading|t[ií]tulo)\s*\d", (p.style.name or "").lower()) is not None
        for r in p.runs:
            n = len(r.text.strip())
            if not n:
                continue
            fu = est.fuente(r, p) or "?"
            ta = est.tamano(r, p)
            fuentes[fu] += n
            total += n
            if es_titulo:
                continue  # los titulos pueden tener otro tamano; solo se les revisa la fuente
            tamanos[ta] += n
            if len(ejemplos_f[fu]) < 3 and fu != f["fuente"]:
                ejemplos_f[fu].append(ubic(i, rangos, bloques, r.text))
            if len(ejemplos_t[ta]) < 3 and ta != f["tamano_pt"]:
                ejemplos_t[ta].append(ubic(i, rangos, bloques, r.text))
    if total:
        for fu, n in fuentes.items():
            if fu.lower() != f["fuente"].lower():
                pct = n / total
                sev = "Error" if pct > f["tolerancia_fuera_de_norma"] else "Advertencia"
                obs.add("Formato", sev, juntar(ejemplos_f[fu]) or "Documento",
                        f"Texto en fuente '{fu}' ({pct:.1%} del documento)", f"La fuente debe ser {f['fuente']}")
        total_t = sum(tamanos.values()) or 1
        for ta, n in tamanos.items():
            if abs(ta - f["tamano_pt"]) > 0.1:
                pct = n / total_t
                sev = "Error" if pct > f["tolerancia_fuera_de_norma"] else "Advertencia"
                obs.add("Formato", sev, juntar(ejemplos_t[ta]) or "Documento",
                        f"Texto en tamaño {ta:g} pt ({pct:.1%} del documento)", f"El tamaño debe ser {f['tamano_pt']} pt")

    # interlineado y alineacion en parrafos de cuerpo (largos, fuera de tablas)
    inter_mal, alin_mal = [], []
    for i, b in enumerate(bloques):
        if not isinstance(b, Paragraph) or len(b.text.strip()) < 120:
            continue
        if abs(est.interlineado(b) - f["interlineado"]) > 0.05:
            inter_mal.append((i, est.interlineado(b), b.text))
        if f["alineacion_cuerpo"] == "justificado" and est.alineacion(b) != WD_ALIGN_PARAGRAPH.JUSTIFY:
            alin_mal.append((i, b.text))
    if inter_mal:
        obs.add("Formato", "Error", juntar([ubic(i, rangos, bloques, t, 30) for i, _, t in inter_mal[:4]]),
                f"Interlineado distinto de {f['interlineado']} en {len(inter_mal)} párrafo(s)",
                "Valores encontrados: " + ", ".join(sorted({str(v) for _, v, _ in inter_mal})))
    if alin_mal:
        obs.add("Formato", "Error", juntar([ubic(i, rangos, bloques, t, 30) for i, t in alin_mal[:4]]),
                f"{len(alin_mal)} párrafo(s) de texto sin justificar")


