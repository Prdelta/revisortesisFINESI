"""
Tests de cómo se lee el modelo de Word, donde python-docx tiene dos trampas que
generaban observaciones falsas sobre documentos correctos:

1. `first_line_indent` del párrafo es None cuando la sangría la define el estilo, y
   Word normalmente la define ahí. Leyendo solo el párrafo, toda lista de referencias
   bien formateada salía como si le faltara la sangría francesa.
2. `row.cells` devuelve la celda combinada una vez por cada columna que abarca, así
   que el texto de un encabezado combinado se contaba dos o tres veces.

Los documentos se arman acá con python-docx: son casos mínimos y exactos, y así no
hace falta versionar un .docx por caso.
"""
import pytest
from docx import Document
from docx.shared import Cm

from core.citas import revisar_citas
from core.extension import contar_palabras
from core.utils import Obs, bloques_en_orden, celdas_unicas, sangria_francesa

REFERENCIAS = [
    "Apaza, L. (2019). Anemia infantil en zonas altoandinas. Revista Peruana, 12(3), 45-58.",
    "Cruz, M. (2020). Determinantes de la desnutrición crónica. Editorial Universitaria.",
    "Quispe, J. (2021). Prevalencia de anemia en escolares. Revista Andina, 8(1), 10-22.",
]

REGLAS = {"citas": {"estilo": "APA7", "seccion_referencias": "Referencias"}}


def _guardar(doc, tmp_path, nombre="doc.docx"):
    ruta = tmp_path / nombre
    doc.save(str(ruta))
    return Document(str(ruta))     # se relee para leer lo que quedó en el archivo


def _observaciones_de_sangria(doc):
    bloques = list(bloques_en_orden(doc))
    obs = Obs()
    revisar_citas(bloques, {"Referencias": (0, len(bloques))}, REGLAS, obs)
    return [x["observacion"] for x in obs.items if "francesa" in x["observacion"]]


# --------------------------------------------------------------------------- sangría francesa

def test_la_sangria_definida_en_el_estilo_cuenta(tmp_path):
    """
    El caso que daba un Error falso: las tres referencias tienen sangría francesa, pero
    puesta en el estilo, que es como la pone Word.
    """
    doc = Document()
    estilo = doc.styles.add_style("Referencias APA", 1)
    estilo.paragraph_format.first_line_indent = Cm(-1.27)
    doc.add_paragraph("Referencias")
    for texto in REFERENCIAS:
        doc.add_paragraph(texto, style="Referencias APA")
    assert _observaciones_de_sangria(_guardar(doc, tmp_path)) == []


def test_la_sangria_definida_en_el_parrafo_cuenta(tmp_path):
    doc = Document()
    doc.add_paragraph("Referencias")
    for texto in REFERENCIAS:
        p = doc.add_paragraph(texto)
        p.paragraph_format.first_line_indent = Cm(-1.27)
    assert _observaciones_de_sangria(_guardar(doc, tmp_path)) == []


def test_la_sangria_heredada_de_un_estilo_base_cuenta(tmp_path):
    """si el tesista deriva su estilo de otro que ya la trae, también vale"""
    doc = Document()
    base = doc.styles.add_style("Base APA", 1)
    base.paragraph_format.first_line_indent = Cm(-1.27)
    propio = doc.styles.add_style("Referencias del tesista", 1)
    propio.base_style = base
    doc.add_paragraph("Referencias")
    for texto in REFERENCIAS:
        doc.add_paragraph(texto, style="Referencias del tesista")
    assert _observaciones_de_sangria(_guardar(doc, tmp_path)) == []


def test_sin_sangria_se_sigue_observando(tmp_path):
    """el arreglo no puede haber apagado la comprobación"""
    doc = Document()
    doc.add_paragraph("Referencias")
    for texto in REFERENCIAS:
        doc.add_paragraph(texto)
    observaciones = _observaciones_de_sangria(_guardar(doc, tmp_path))
    assert observaciones == ["3 de 3 referencias sin sangría francesa"]


def test_una_sangria_positiva_no_es_francesa(tmp_path):
    """sangría de primera línea hacia adentro es lo contrario de la francesa"""
    doc = Document()
    doc.add_paragraph("Referencias")
    for texto in REFERENCIAS:
        p = doc.add_paragraph(texto)
        p.paragraph_format.first_line_indent = Cm(1.27)
    assert _observaciones_de_sangria(_guardar(doc, tmp_path)) != []


def test_sangria_francesa_sobre_un_parrafo_suelto(tmp_path):
    doc = Document()
    con = doc.add_paragraph("con")
    con.paragraph_format.first_line_indent = Cm(-1.27)
    doc.add_paragraph("sin")
    doc = _guardar(doc, tmp_path)
    assert sangria_francesa(doc.paragraphs[0]) is True
    assert sangria_francesa(doc.paragraphs[1]) is False


# --------------------------------------------------------------------------- celdas combinadas

def _tabla_de_presupuesto(doc):
    """un encabezado combinado a tres columnas, como en todo cronograma o presupuesto"""
    t = doc.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 2)).text = "PRESUPUESTO GENERAL DEL PROYECTO"   # 4 palabras
    t.cell(1, 0).text = "Materiales"                                            # 1
    t.cell(1, 1).text = "Servicios"                                             # 1
    t.cell(1, 2).text = "Total"                                                 # 1
    return t


def test_la_celda_combinada_se_cuenta_una_vez(tmp_path):
    """
    Contándola una vez por columna, el total se inflaba y la comparación con las palabras
    que declara Turnitin acusaba en falso al tesista de haber pasado otro documento.
    """
    doc = Document()
    doc.add_paragraph("Uno dos tres")      # 3
    _tabla_de_presupuesto(doc)             # 4 + 1 + 1 + 1
    assert contar_palabras(list(bloques_en_orden(_guardar(doc, tmp_path)))) == 10


def test_celdas_unicas_no_repite_la_combinada(tmp_path):
    doc = Document()
    _tabla_de_presupuesto(doc)
    tabla = _guardar(doc, tmp_path).tables[0]
    repetidas = [c.text for f in tabla.rows for c in f.cells]
    unicas = [c.text for c in celdas_unicas(tabla)]
    assert len(repetidas) == 6, "row.cells repite la combinada: es la trampa que se evita"
    assert unicas == ["PRESUPUESTO GENERAL DEL PROYECTO", "Materiales", "Servicios", "Total"]


def test_una_tabla_sin_combinar_no_cambia(tmp_path):
    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    for celda, texto in zip([t.cell(0, 0), t.cell(0, 1), t.cell(1, 0), t.cell(1, 1)],
                            ["uno", "dos", "tres", "cuatro"]):
        celda.text = texto
    doc = _guardar(doc, tmp_path)
    assert contar_palabras(list(bloques_en_orden(doc))) == 4
    assert len(list(celdas_unicas(doc.tables[0]))) == 4


def test_una_combinacion_vertical_tampoco_se_repite(tmp_path):
    doc = Document()
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).merge(t.cell(2, 0)).text = "Etapa uno"       # 2 palabras
    for fila, texto in zip(range(3), ["enero", "febrero", "marzo"]):
        t.cell(fila, 1).text = texto                          # 1 cada una
    doc = _guardar(doc, tmp_path)
    assert contar_palabras(list(bloques_en_orden(doc))) == 5


@pytest.mark.parametrize("palabras_informe, deberia_observar", [
    (10, False),    # coincide con el .docx
    (30, True),     # el informe es de otro documento
])
def test_el_conteo_arreglado_no_dispara_la_alarma_de_turnitin(tmp_path, palabras_informe,
                                                              deberia_observar):
    """
    El efecto de punta a punta: con el conteo inflado, un documento con una tabla
    combinada disparaba 'el documento que se pasó por Turnitin no coincide'.
    """
    from core.turnitin import Informe, revisar_similitud

    doc = Document()
    doc.add_paragraph("Uno dos tres")
    _tabla_de_presupuesto(doc)
    palabras_docx = contar_palabras(list(bloques_en_orden(_guardar(doc, tmp_path))))

    obs = Obs()
    informe = Informe(ruta="t.pdf", archivo="t.pdf", es_turnitin=True, similitud=5.0,
                      palabras=palabras_informe)
    revisar_similitud(informe, {"similitud": {"max_pct": None}}, obs,
                      palabras_docx=palabras_docx, nombre_docx="tesis.docx")
    hay = any("no coincide" in x["observacion"] for x in obs.items)
    assert hay is deberia_observar
