"""
Piezas compartidas por los tests.

Los tests corren sobre los documentos de 'ejemplos/' y no necesitan Java ni
LibreOffice: la ortografía se fuerza al diccionario hunspell que viene en 'data/dic',
así el resultado es el mismo en cualquier máquina. La página y línea de cada
observación sí dependen de LibreOffice, por eso las instantáneas no las incluyen.
"""
import os
import sys

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "src"))

EJEMPLOS = os.path.join(RAIZ, "ejemplos")
INSTANTANEAS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "instantaneas")


@pytest.fixture(scope="session")
def ejemplos():
    return EJEMPLOS


@pytest.fixture(autouse=True)
def datos_aislados(tmp_path, monkeypatch):
    """
    Aparta los dos archivos que el revisor edita a mano — 'mis_reglas.yaml' y
    'registro.csv' — hacia una carpeta temporal.

    Sin esto, las instantáneas dependerían de las reglas propias de quien corre los
    tests y se romperían cada vez que alguien agrega una. Las reglas del programa
    (data/reglas/) sí se usan de verdad: son parte de lo que se está probando.
    """
    propios = tmp_path / "datos"
    propios.mkdir()
    monkeypatch.setattr("core.orquestador.DATOS", str(propios))
    return str(propios)


@pytest.fixture(autouse=True)
def sin_word(monkeypatch):
    """
    Word ubica la línea de cada observación, pero abrirlo en cada test tarda ~10 s por
    documento y hace depender los tests de que esté instalado. La página y línea ya
    quedan fuera de las instantáneas por la misma razón (ver observaciones_del_excel).
    """
    monkeypatch.setattr("utilidades.localizador.hay_word", lambda: False)


@pytest.fixture
def salida(tmp_path):
    """carpeta de reportes propia de cada test"""
    d = tmp_path / "reportes"
    d.mkdir()
    return str(d)


@pytest.fixture
def opciones(salida):
    """
    Opciones de una corrida reproducible: sin LanguageTool (que descarga 260 MB y
    necesita Java) y sin los datos del expediente, que cambiarían el reporte.
    """
    from core.orquestador import Opciones
    return Opciones(tipo="proyecto", salida=salida, excel=True, motor="diccionario",
                    revisor="Revisor de prueba", formato="resumen")


def comparar(nombre, actual):
    """
    Compara con la instantánea congelada. Si el comportamiento cambió a propósito:
        $env:REGENERAR_INSTANTANEAS="1"; python -m pytest tests/
    """
    ruta = os.path.join(INSTANTANEAS, nombre + ".txt")
    if os.environ.get("REGENERAR_INSTANTANEAS"):
        os.makedirs(INSTANTANEAS, exist_ok=True)
        with open(ruta, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(actual)
        pytest.skip(f"instantánea regenerada: {nombre}")
    if not os.path.exists(ruta):
        pytest.fail(f"falta la instantánea '{nombre}'. Generarla con:\n"
                    f'  $env:REGENERAR_INSTANTANEAS="1"; python -m pytest tests/')
    with open(ruta, encoding="utf-8") as fh:
        esperado = fh.read()
    assert actual == esperado, (
        f"el motor ya no observa lo mismo en '{nombre}'.\n"
        "Si el cambio es a propósito, regenerar la instantánea; si no, es una regresión.")


def observaciones_del_excel(ruta_xlsx):
    """
    Lee el Excel del reporte y devuelve una línea por observación. Se usa el Excel
    porque es la salida que lista los hallazgos uno por uno. La columna de página y
    línea se omite a propósito: depende de LibreOffice y no está en todas las máquinas.
    """
    from openpyxl import load_workbook
    ws = load_workbook(ruta_xlsx)["Observaciones"]
    filas = []
    for fila in ws.iter_rows(min_row=5, values_only=True):
        if not fila or fila[1] is None:
            continue
        _, categoria, severidad, _linea, ubicacion, observacion = fila[:6]
        filas.append(" | ".join(str(x or "") for x in (categoria, severidad, ubicacion, observacion)))
    return filas
